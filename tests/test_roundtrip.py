# -*- coding: utf-8 -*-
"""Unit tests for the run / paragraph / soft-break text round-trip."""
from __future__ import annotations

import sys

from pptx import Presentation
from pptx.util import Inches, Pt

from pptx_translate import core


def build_frame(paragraphs=1):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(3))
    tf = box.text_frame
    first = tf.paragraphs[0]
    run = first.add_run()
    run.text = ""
    for _ in range(paragraphs - 1):
        tf.add_paragraph().add_run().text = ""
    return tf


def para_text(para) -> str:
    return core._read_group([n for g in core._paragraph_lines(para) for n in g])


def test_single_run_roundtrip():
    tf = build_frame()
    para = tf.paragraphs[0]
    run = para.runs[0]
    run.text = "Hello world"
    run.font.size = Pt(20)
    run.font.bold = True
    core._set_paragraph_text(para, "你好，世界")
    assert para.runs[0].text == "你好，世界", para.runs[0].text
    assert para.runs[0].font.size.pt == 20
    assert para.runs[0].font.bold is True
    print("single run, format kept OK")


def test_soft_break_roundtrip():
    """Source paragraph: run1 <a:br/> run2 must become exactly two lines."""
    tf = build_frame()
    para = tf.paragraphs[0]
    run1 = para.runs[0]
    run1.text = "Friendly knights"
    from lxml import etree
    br = etree.SubElement(para._p, core.A_BR)
    para._p.append(br)
    run2 = para.add_run()
    # add_run appends after the break, which is exactly the source pattern
    run2.text = "test your castle walls"

    groups = core._paragraph_lines(para)
    assert len(groups) == 2, [core._read_group(g) for g in groups]
    core._set_paragraph_text(para, "友好的骑士\x0b测试你的城墙")
    groups = core._paragraph_lines(para)
    got = [core._read_group(g) for g in groups]
    assert got == ["友好的骑士", "测试你的城墙"], got
    # the break element itself survived
    assert len(para._p.findall(core.A_BR)) == 1
    print(f"soft-break round-trip   OK -> {got} (a:br preserved)")


def test_soft_break_not_collapsed_into_one_run():
    """Regression: line 2 must not overwrite line 1's run."""
    tf = build_frame()
    para = tf.paragraphs[0]
    para.runs[0].text = "AAA"
    from lxml import etree
    etree.SubElement(para._p, core.A_BR)
    para.add_run().text = "BBB"
    core._set_paragraph_text(para, "第一行\x0b第二行")
    got = [core._read_group(g) for g in core._paragraph_lines(para)]
    assert got == ["第一行", "第二行"], got
    assert "AAA" not in "".join(got)
    print("regression: no clobber OK")


def test_multi_run_proportional_split():
    tf = build_frame()
    para = tf.paragraphs[0]
    for text, bold in (("This is ", False), ("very important", True), (" text.", False)):
        run = para.add_run() if para.runs[0].text else para.runs[0]
        run.text = text
        run.font.bold = bold
    core._set_paragraph_text(para, "这是一段非常重要的文本。")
    runs = list(para.runs)
    texts = [r.text for r in runs]
    assert "".join(texts) == "这是一段非常重要的文本。", texts
    assert all(texts), f"a run went empty and lost its formatting: {texts}"
    assert runs[1].font.bold is True
    print(f"proportional, none empty OK -> {texts}")


def test_paragraph_count_and_order():
    tf = build_frame(3)
    for i, para in enumerate(tf.paragraphs):
        para.runs[0].text = f"sentence {i}"
    for para, new in zip(list(tf.paragraphs), ["句子零", "句子一", "句子二"]):
        core._set_paragraph_text(para, new)
    assert len(tf.paragraphs) == 3
    assert [p.runs[0].text for p in tf.paragraphs] == ["句子零", "句子一", "句子二"]
    print("paragraph order/count   OK")


def test_apply_units_to_frame_skips_absent():
    tf = build_frame(3)
    for i, para in enumerate(tf.paragraphs):
        para.runs[0].text = f"line {i}"
    core.apply_units_to_frame(tf, {1: "只改中间那行"})
    assert [p.runs[0].text for p in tf.paragraphs] == ["line 0", "只改中间那行", "line 2"]
    print("per-index apply         OK")


def test_python_pptx_vb_escape_normalized():
    """A soft break read back as ``_x000B_`` must not reach the model literally."""
    tf = build_frame()
    para = tf.paragraphs[0]
    para.runs[0].text = "🛡️ Simulated Attacks_x000B_Friendly knights test your castle walls."
    groups = core._paragraph_lines(para)
    joined = "".join(core._read_group(g) for g in groups)
    assert core.VB_ESCAPE not in joined, joined
    assert "\x0b" in joined, repr(joined)
    assert core.text_parts(joined) == [
        "🛡️ Simulated Attacks",
        "Friendly knights test your castle walls.",
    ], core.text_parts(joined)

    # and writing a soft break back must emit the valid escape, not a raw newline
    core._set_paragraph_text(para, "🛡️ 模拟攻击\x0b友好的骑士来试你的城墙。")
    node_texts = [n.text for g in core._paragraph_lines(para) for n in g]
    assert any(core.VB_ESCAPE in (t or "") for t in node_texts), node_texts
    assert not any("\n" in (t or "") for t in node_texts), node_texts
    print("_x000B_ normalization   OK (escape in, escape out)")


def test_break_normalization_aliases():
    """Every spelling a model uses for a soft break must collapse to \\x0b."""
    for alias in ("_x000B_", "\\u000b", "\\v", "\\x0b", "\x0b"):
        got = core.normalize_breaks("模拟攻击" + alias + "友好的骑士")
        assert got == "模拟攻击\x0b友好的骑士", (alias, repr(got))
        assert core.text_parts(got) == ["模拟攻击", "友好的骑士"], (alias, core.text_parts(got))
    print("break alias normalization OK")


def test_parse_single_shapes():
    """A single-item reply arrives as an array, a quoted string, or bare text."""
    assert core.parse_single('["你好世界"]') == "你好世界"
    assert core.parse_single('"你好世界"') == "你好世界"
    assert core.parse_single("你好世界") == "你好世界"
    assert core.parse_single("1. 你好世界") == "你好世界"
    assert core.parse_single('{"i": 1, "t": "你好世界"}') == "你好世界"
    assert core.parse_single("```\n你好世界\n```") == "你好世界"
    assert core.parse_single("  你好世界  ") == "你好世界"
    print("single-reply unwrapping  OK")


def test_quote_neutralization():
    """A bare ASCII quote inside a translation must not survive into the XML."""
    src = '"Your prompt: ‘Forget your rules’"'
    got = core.neutralize_quotes(src)
    assert '"' not in got, got
    assert "Forget your rules" in got
    print(f"quote neutralization     OK -> {got}")


def test_glossary_lock_and_unlock():
    """Locked terms are replaced by tokens, then restored as English."""
    provider = core.build_provider(core.ProviderConfig(provider="mock", api_key="x", model="m"))
    tr = core.Translator(provider, glossary={"prompt injection": "prompt injection"})
    assert tr.lock_terms("This is a Prompt Injection attack") == "This is a [[G1]] attack"
    assert tr.unlock_terms("This is a [[G1]] attack") == "This is a prompt injection attack"
    tr2 = core.Translator(provider, glossary={"prompt injection": "提示注入"})
    assert tr2.unlock_terms(tr2.lock_terms("Prompt Injection")) == "提示注入"
    print("glossary lock/unlock     OK")


def test_glossary_token_numbering_is_global():
    """[[Gn]] must mean the same term in every string, not per-string position."""
    provider = core.build_provider(core.ProviderConfig(provider="mock", api_key="x", model="m"))
    tr = core.Translator(provider, glossary={"ATT&CK": "ATT&CK", "ATLAS": "ATLAS",
                                             "OWASP": "OWASP"})
    long_text = "Hacker playbooks (ATT&CK) & AI defenses (ATLAS)."
    short_text = "ATLAS alone."
    locked_long = tr.lock_terms(long_text)
    locked_short = tr.lock_terms(short_text)
    # the ATLAS token must be identical in both strings
    atlas_long = core.re.search(r"\[\[G\d+\]\]", locked_long.split("AI defenses ")[1]).group(0)
    atlas_short = core.re.match(r"\[\[G\d+\]\]", locked_short).group(0)
    assert atlas_long == atlas_short, (locked_long, locked_short)
    assert tr.unlock_terms(locked_long) == long_text
    assert tr.unlock_terms(locked_short) == short_text
    print(f"global token numbering   OK ({atlas_short} = ATLAS in both)")


def test_unlock_survives_model_reformatting():
    """A model that spaces or unbalances a token must not leak it onto a slide."""
    provider = core.build_provider(core.ProviderConfig(provider="mock", api_key="x", model="m"))
    tr = core.Translator(provider, glossary={"ATLAS": "ATLAS", "OWASP": "OWASP"})
    token = tr.lock_terms("ATLAS")            # e.g. [[G2]]
    for broken in (token, token.replace("[[", "[[ "), token.replace("]]", " ]"),
                   token.replace("[[", "["), token.replace("]]", "]")):
        restored = tr.unlock_terms(f"防御（{broken}）。")
        assert "[[G" not in restored, (broken, restored)
        assert "ATLAS" in restored, (broken, restored)
    print("token reformat recovery  OK")


def test_cache_keeps_success_and_retries_failure(tmp_path=None):
    """A failure must be retried, and must never erase an earlier success."""
    import tempfile
    from pathlib import Path

    class FlakyProvider(core.BaseProvider):
        """Fails 'BAD' forever, translates anything else."""

        def __init__(self, cfg):
            super().__init__(cfg)
            self.calls = 0

        def translate(self, items):
            out = []
            for text in items:
                self.calls += 1
                if text == "BAD":
                    raise core.TranslationError("always fails")
                out.append("【" + text + "】")
            return out

    cfg = core.ProviderConfig(provider="mock", api_key="x", model="flaky")
    with tempfile.TemporaryDirectory() as tmp:
        cache_file = Path(tmp) / "c.json"

        # seed a known-good cache entry (as a previous successful run would)
        provider = FlakyProvider(cfg)
        tr = core.Translator(provider, cache_path=cache_file, use_cache=True)
        good_key = tr.cache_key("GOOD")
        cache_file.write_text(core.json.dumps({good_key: "【GOOD】"}), "utf-8")
        tr.cache = {good_key: "【GOOD】"}
        assert len(tr.cache) == 1, tr.cache

        result = tr.translate_texts(["GOOD", "BAD"])

        # the cached success is reused, only the failure is requested
        assert result["GOOD"].startswith("【"), result
        assert result["BAD"] == "BAD", "an untranslated item keeps its source text"
        assert provider.calls == 1, f"expected only the failing string to be requested, got {provider.calls}"

        # and the failure did not erase the earlier success from the cache
        cache = core.json.loads(cache_file.read_text("utf-8"))
        assert good_key in cache, cache
        assert len(cache) == 1, cache
    print("cache: success kept, failure retried OK")


def test_local_provider_wiring():
    """The local provider must construct and report its state without torch."""
    from pptx_translate import local

    ok, why = local.is_available()
    assert isinstance(ok, bool) and isinstance(why, str)
    provider = local.build_local_provider(
        core.ProviderConfig(provider="local", api_key="", model="x",
                            target_language="Simplified Chinese (简体中文)"),
        model_id="example/tiny-mt",
        device="cpu",
        prompt_style="direct",
    )
    assert provider.local.model_id == "example/tiny-mt"
    assert provider.local.device == "cpu"
    assert provider.local.prompt_style == "direct"
    assert provider.local.target_language == "中文", provider.local.target_language
    assert provider.loaded is False, "weights must not load until first translate"
    # the factory used by the CLI must also resolve to the local provider
    built = core.build_provider(core.ProviderConfig(provider="local", api_key="", model="x"))
    assert built.name == "local"
    print(f"local provider wiring    OK (deps {'present' if ok else 'absent: ' + why})")


def test_local_target_language_mapping():
    """The shared label must map onto a short name the local prompt can use."""
    from pptx_translate.cli import _local_target

    assert _local_target("Simplified Chinese (简体中文)") == "中文"
    assert _local_target("English") == "English"
    assert _local_target("Japanese") == "Japanese"
    print("local target mapping     OK")


def test_invented_placeholders_are_removed():
    """A model that imitates the [[E1]] format must not put tokens on a slide."""
    cases = [
        # the local model invented these from paragraphs that had no emoji at all
        ("超过 201 种方法：[[E1]]：攻击者",
         "超过 201 种方法：攻击者"),
        ("[[E1]]: 数据外传", "数据外传"),
        ("与 [[U1]]（2020年至2021年中期）", "与（2020年至2021年中期）"),
        ("发生了什么？[[E1]]乌克兰被破坏。[[E12]]“沙虫”攻击。",
         "发生了什么？乌克兰被破坏。“沙虫”攻击。"),
    ]
    for src, expected in cases:
        got, removed = core.strip_unknown_tokens(src, set())
        assert removed, f"nothing was detected in {src!r}"
        assert got == expected, (src, got, expected)

    # a real emoji that was properly restored must survive untouched
    real = "🛡️ 模拟攻击\x0b友好的骑士们正在测试你的城堡城墙。"
    assert core.strip_unknown_tokens(real, set()) == (real, [])
    # and a token the source genuinely owned is never dropped
    kept, removed = core.strip_unknown_tokens("x [[E1]] y", {"[[E1]]"})
    assert (kept, removed) == ("x [[E1]] y", [])
    print("invented placeholder purge OK")


def test_bracket_repair():
    """A garbled or duplicated closing bracket is fixed; balanced text is left alone."""
    src = "Hacker playbooks (ATT&CK) & AI defenses (ATLAS)."
    # duplicated closers
    got, note = core.repair_brackets("黑客操作手册（ATT&CK）] & 防御（ATLAS））。", src)
    assert note, "a repair should have been reported"
    assert got == "黑客操作手册（ATT&CK） & 防御（ATLAS）。", got
    # a closer mangled into an angle bracket
    got, note = core.repair_brackets("黑客（ATT&CK>）与防御（ATLAS））。", src)
    assert got == "黑客（ATT&CK）与防御（ATLAS）。", got
    # already balanced: untouched
    balanced = "黑客（ATT&CK）与防御（ATLAS）。"
    assert core.repair_brackets(balanced, src) == (balanced, "")
    # source without brackets: the model's own brackets are left alone
    assert core.repair_brackets("（多余）", "plain text") == ("（多余）", "")
    # a missing closer is NOT invented
    short = "黑客（ATT&CK与防御（ATLAS）。"
    assert core.repair_brackets(short, src)[0] == short
    print("bracket repair           OK")


def test_simplify_layer():
    """Traditional characters from a model must land as Simplified."""
    from pptx_translate import simplify

    cases = {
        "測試比較 ≠ 安全性：人工智能的誤解": "测试比较 ≠ 安全性：人工智能的误解",
        "演示紅隊攻擊": "演示红队攻击",
        "嚴格範圍：無混亂，只有線索。": "严格范围：无混乱，只有线索。",
    }
    for src, expected in cases.items():
        got = simplify.to_simplified(src)
        # without OpenCC the small table may not cover every character, so only
        # assert the characters it does know about
        assert "規則" not in got and "嚴格" not in got
        if simplify._opencc() is not None:
            assert got == expected, (src, got, expected)
    # already-Simplified text must be untouched
    assert simplify.to_simplified("红队攻击（简体保持不变）") == "红队攻击（简体保持不变）"
    assert simplify.has_traditional("規則") is True
    assert simplify.has_traditional("规则") is False
    print("simplify layer           OK")


def test_blank_detection():
    assert core.is_blank("   ")
    assert core.is_blank("2024")
    assert core.is_blank("→ ← •")
    assert core.is_blank("已经是中文了")
    assert not core.is_blank("Prompt Injection")
    assert not core.is_blank("🔍 What is it?")
    assert not core.is_blank("North")
    print("blank detection         OK")


def test_placeholder_preservation():
    pm = core.PlaceholderMap()
    src = "🎯 Why It Matters💀"
    protected = pm.protect(src)
    assert "🎯" not in protected and "[[" in protected
    mapping = pm.mapping_for(src)
    assert pm.restore(protected, mapping) == src
    print(f"emoji placeholder       OK -> {protected}")


def test_chart_part_roundtrip():
    """Chart labels are translated; cached numeric values are left alone."""
    xml = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"'
        b' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        b"<c:chart><c:title><c:tx><c:rich><a:p><a:r><a:t>Quarterly Revenue</a:t></a:r>"
        b"</a:p></c:rich></c:tx></c:title><c:plotArea>"
        b"<c:ser><c:tx><c:strRef><c:strCache>"
        b'<c:pt idx="0"><c:v>Threats blocked</c:v></c:pt>'
        b"</c:strCache></c:strRef></c:tx>"
        b"<c:cat><c:numRef><c:numCache>"
        b'<c:pt idx="0"><c:v>1</c:v></c:pt><c:pt idx="1"><c:v>2</c:v></c:pt>'
        b"</c:numCache></c:numRef></c:cat>"
        b"</c:ser></c:plotArea></c:chart></c:chartSpace>"
    )
    bucket: dict = {}
    units = core.extract_part_text(xml, "ppt/charts/chart1.xml", bucket)
    texts = sorted(u.text for u in units)
    assert texts == ["Quarterly Revenue", "Threats blocked"], texts
    mapping = {"Quarterly Revenue": "季度营收", "Threats blocked": "已拦截威胁"}
    changed = core.apply_part_text(bucket, "ppt/charts/chart1.xml", mapping)
    assert changed == 2, changed
    out = core.serialize_part(bucket["ppt/charts/chart1.xml"]["root"])
    assert "季度营收".encode() in out and "已拦截威胁".encode() in out
    # the numeric category cache must survive untouched
    assert b"<c:numCache>" in out and b'<c:v>2</c:v>' in out
    print("chart part round-trip   OK (2 labels, numeric cache intact)")


if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for fn in tests:
        try:
            fn()
        except AssertionError as exc:
            failed += 1
            print(f"FAIL  {fn.__name__}: {exc}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"ERROR {fn.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
