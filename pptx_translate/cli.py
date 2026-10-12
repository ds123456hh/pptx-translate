# -*- coding: utf-8 -*-
"""pptx_translate.cli — command line entry point."""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import tempfile
import time
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from . import core, fit, local

log = logging.getLogger("pptx_translate")


def _local_target(target_language: str) -> str:
    """Map the shared target-language label onto a short name for the local model."""
    text = (target_language or "").lower()
    if "chinese" in text or "中文" in text:
        return "中文"
    if "english" in text or "英文" in text:
        return "English"
    return target_language or "中文"

DEFAULT_BASE_URL = "https://api.deepseek.com/v1"
DEFAULT_MODEL = "deepseek-chat"
DEFAULT_KEY_ENV = "DEEPSEEK_API_KEY"
# Local Ollama server. `ollama create hy-mt-local -f Modelfile` registers the tag.
DEFAULT_OLLAMA_HOST = "http://localhost:11434/v1"
DEFAULT_OLLAMA_MODEL = "hy-mt-local"
# Checked in order; the first one set wins.
KEY_ENV_CANDIDATES = (
    "DSH_TRANSLATE_API_KEY",
    "DEEPSEEK_API_KEY",
    "TOKENHUB_API_KEY",
    "SILICONFLOW_API_KEY",
    "OPENAI_API_KEY",
)


def resolve_api_key(explicit: str | None) -> tuple[str | None, str | None]:
    """Find a key from --api-key or the environment. Returns (key, source)."""
    if explicit:
        return explicit, "--api-key"
    for name in KEY_ENV_CANDIDATES:
        value = os.environ.get(name)
        if value:
            return value, name
    return None, None


# --------------------------------------------------------------------------- #
# argument parsing
# --------------------------------------------------------------------------- #
def _safe_print(text: str) -> None:
    """Print text that may contain emoji, on a console that cannot encode it.

    A Windows console defaults to a legacy code page (GBK here), and printing a
    preview containing an emoji raises UnicodeEncodeError and kills the command
    after the work is already done. Fall back to UTF-8 bytes on the underlying
    buffer, and only drop characters as a last resort.
    """
    try:
        print(text)
        return
    except UnicodeEncodeError:
        pass
    try:
        buffer = getattr(sys.stdout, "buffer", None)
        if buffer is not None:
            buffer.write(text.encode("utf-8", "replace"))
            buffer.write(b"\n")
            buffer.flush()
            return
    except Exception:
        pass
    try:
        print(text.encode("ascii", "replace").decode("ascii"))
    except Exception:
        pass


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="pptx_translate",
        description="Translate a .pptx into Chinese in place, preserving every layout, "
                    "style, shape, animation and image byte-for-byte.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "examples:\n"
            "  python -m pptx_translate deck.pptx                 # -> deck.zh-CN.pptx\n"
            "  python -m pptx_translate deck.pptx -o out.pptx --provider deepseek\n"
            "  python -m pptx_translate deck.pptx --dry-run        # inventory, no API calls\n"
            "  python -m pptx_translate deck.pptx --provider mock  # offline pipeline test\n"
        ),
    )
    p.add_argument("input", help="source .pptx")
    p.add_argument("-o", "--output", help="destination .pptx (default: <name>.zh-CN.pptx)")
    p.add_argument("--provider", default="deepseek",
                   choices=["ollama", "local", "deepseek", "openai", "mock"],
                   help="translation backend: 'ollama' talks to a local Ollama server "
                        "(no key, no cost), 'local' runs a HuggingFace model in-process, "
                        "'deepseek'/'openai' use a remote API, 'mock' is offline")
    p.add_argument("--ollama-model", default=DEFAULT_OLLAMA_MODEL,
                   help=f"Ollama model name (default: {DEFAULT_OLLAMA_MODEL})")
    p.add_argument("--ollama-host", default=DEFAULT_OLLAMA_HOST,
                   help=f"Ollama base url (default: {DEFAULT_OLLAMA_HOST})")
    p.add_argument("--local-model", default=None,
                   help="HuggingFace repo for --provider local "
                        "(default: {}".format(core.DEFAULT_LOCAL_MODEL) + ")")
    p.add_argument("--local-device", default="auto", choices=["auto", "cpu", "cuda"],
                   help="device for --provider local (default: auto)")
    p.add_argument("--local-prompt-style", default="auto",
                   choices=["auto", "instruction", "chat", "direct"],
                   help="how the local model is prompted (default: auto)")
    p.add_argument("--local-cache-dir", default=None,
                   help="where downloaded weights are cached (default: HF_HOME)")
    p.add_argument("--local-max-new-tokens", type=int, default=400,
                   help="generation cap per paragraph for --provider local")
    p.add_argument("--local-files-only", action="store_true",
                   help="refuse to download; use only already-cached weights")
    p.add_argument("--base-url", default=None, help=f"API base url (default: {DEFAULT_BASE_URL})")
    p.add_argument("--api-key", default=None,
                   help=f"API key; prefer the {DEFAULT_KEY_ENV} environment variable")
    p.add_argument("--model", default=None, help=f"model id (default: {DEFAULT_MODEL})")
    p.add_argument("--target-language", default="Simplified Chinese (简体中文)")
    p.add_argument("--batch-size", type=int, default=1,
                   help="strings per API call (default 1). MT models merge adjacent "
                        "segments, which makes any batched reply unsafe to map back")
    p.add_argument("--concurrency", type=int, default=6, help="parallel API calls (default 6)")
    p.add_argument("--timeout", type=float, default=180.0, help="per-request timeout seconds")
    p.add_argument("--retries", type=int, default=4, help="attempts per batch")
    p.add_argument("--glossary", default=None,
                   help="JSON file or CSV 'source,target' lines; overrides the built-in "
                        "proper-noun list (e.g. '\"prompt injection,提示注入\"')")
    p.add_argument("--no-glossary", action="store_true",
                   help="do not lock any term to English")
    p.add_argument("--strings-json", default=None,
                   help="pre-translated {'source': 'translation'} JSON; skips the API entirely "
                        "(use for human-reviewed or offline translations)")
    p.add_argument("--keep-terms", default=core.DEFAULT_KEEP_TERMS,
                   help="comma-separated terms the model must leave in English")
    p.add_argument("--cache", default=None,
                   help="translation memory JSON (default: <output>.cache.json)")
    p.add_argument("--no-cache", action="store_true", help="ignore and do not write the cache")
    p.add_argument("--autofit", default="off", choices=["off", "all", "no-titles"],
                   help="shrink oversized text: off (default, fonts untouched), "
                        "no-titles (keep title placeholders at their original size), all")
    p.add_argument("--autofit-preview", action="store_true",
                   help="report which shapes overflow after translation, without changing fonts")
    p.add_argument("--notes", action="store_true", help="also translate speaker notes")
    p.add_argument("--dry-run", action="store_true", help="list what would be translated, then stop")
    p.add_argument("--report", default=None, help="write a JSON report here")
    p.add_argument("--dump-strings", default=None, help="write all source strings to a JSON file")
    p.add_argument("--verbose", "-v", action="count", default=0)
    p.add_argument("--quiet", "-q", action="store_true")
    return p


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def is_local_endpoint(url: str) -> bool:
    """True for a local inference server (Ollama, LM Studio, llama.cpp server).

    Such servers need no credentials, so requiring a key would be wrong.
    """
    text = (url or "").lower()
    return any(
        marker in text
        for marker in ("localhost", "127.0.0.1", "0.0.0.0", "[::1]", "host.docker.internal")
    )


def setup_logging(verbose: int, quiet: bool) -> None:
    level = logging.WARNING if quiet else (logging.DEBUG if verbose else logging.INFO)
    logging.basicConfig(
        level=level,
        format="%(levelname)-7s %(message)s",
        stream=sys.stdout,
    )


def default_output(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}.zh-CN{input_path.suffix}")


def _sidecar(path: Path, suffix: str) -> Path:
    """``deck.pptx`` + ``.report.json`` -> ``deck.pptx.report.json``."""
    return Path(str(path) + suffix)


def load_glossary(spec: str | None) -> dict[str, str]:
    if not spec:
        return {}
    path = Path(spec)
    if not path.exists():
        raise SystemExit(f"glossary file not found: {path}")
    text = path.read_text("utf-8")
    if path.suffix.lower() == ".json":
        data = json.loads(text)
        if not isinstance(data, dict):
            raise SystemExit("glossary JSON must be an object of source -> target")
        return {str(k): str(v) for k, v in data.items()}
    out: dict[str, str] = {}
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        parts = re.split(r"[,\t]", line, maxsplit=1)
        if len(parts) == 2:
            out[parts[0].strip()] = parts[1].strip()
    return out


def visible_parts(src: Path) -> list[str]:
    """Chart / diagram data parts present in the package."""
    found: list[str] = []
    with zipfile.ZipFile(src) as zf:
        for name in zf.namelist():
            if name.startswith("ppt/charts/chart") or name.startswith("ppt/diagrams/data"):
                if name.endswith(".xml"):
                    found.append(name)
    return sorted(found)


def extract_part_units(src: Path, buckets: dict, skip: set[str]) -> list[core.Unit]:
    """Collect text from chart and diagram data parts."""
    units: list[core.Unit] = []
    with zipfile.ZipFile(src) as zf:
        names = [n for n in zf.namelist()
                 if n.endswith(".xml") and (n.startswith("ppt/charts/chart") or n.startswith("ppt/diagrams/data"))]
        for name in sorted(names):
            xml = zf.read(name)
            for unit in core.extract_part_text(xml, name, buckets):
                if unit.source in skip or core.is_blank(unit.source):
                    continue
                units.append(unit)
    return units


# --------------------------------------------------------------------------- #
# apply translations to a presentation object
# --------------------------------------------------------------------------- #
def apply_unitwise(prs, units: list[core.Unit], resolved: dict[str, str], translate_notes: bool, stats: dict) -> None:
    """Replay the exact collection walk and write each translation into its paragraph.

    Every unit remembers the frame it came from (``where``) and the paragraph
    index inside that frame, so blank-enough-to-skip paragraphs can never shift
    a translation onto a neighbouring line.
    """
    by_slide: dict[int, list[core.Unit]] = {}
    for unit in units:
        by_slide.setdefault(unit.slide, []).append(unit)
    notes_units = [u for u in units if u.kind == "notes"]
    if not translate_notes and notes_units:
        log.info("skipping %d speaker-note paragraph(s)", len(notes_units))

    for slide_no, slide in enumerate(prs.slides, 1):
        todo = [u for u in by_slide.get(slide_no, []) if u.kind != "notes" or translate_notes]
        if not todo:
            continue
        grouped: dict[tuple[str, str], list[core.Unit]] = {}
        for unit in todo:
            grouped.setdefault((unit.where, unit.kind), []).append(unit)
        frames: list[tuple[object, tuple[str, str]]] = []

        def sink(shape, s_no: int, where: str) -> None:
            label = f"{where}/{shape.name}"
            if shape.has_text_frame:
                frames.append((shape.text_frame, (label, "shape")))
            if getattr(shape, "has_table", False) and shape.has_table:
                for r, row in enumerate(shape.table.rows):
                    for c, cell in enumerate(row.cells):
                        frames.append((cell.text_frame, (f"{label}/table[{r}][{c}]", "table")))

        core._walk_shapes(slide.shapes, slide_no, f"slide{slide_no}", sink)
        if translate_notes and slide.has_notes_slide:
            frames.append((slide.notes_slide.notes_text_frame, (f"slide{slide_no}/notes", "notes")))

        for tf, key in frames:
            group = grouped.get(key)
            if not group:
                continue
            by_index: dict[int, str] = {}
            for unit in group:
                new = resolved.get(unit.source)
                if new and new != unit.source:
                    by_index[unit.index] = new
            if by_index and core.apply_units_to_frame(tf, by_index):
                stats["frames"] += 1


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    setup_logging(args.verbose, args.quiet)
    started = time.time()

    src = Path(args.input).expanduser()
    if not src.exists():
        log.error("input not found: %s", src)
        return 2
    if src.suffix.lower() not in (".pptx", ".pptm"):
        log.error("input must be .pptx or .pptm (got %s)", src.suffix)
        return 2

    out = Path(args.output).expanduser() if args.output else default_output(src)
    out = out if out.is_absolute() else (Path.cwd() / out)
    cache_path = None if args.no_cache else (Path(args.cache) if args.cache else _sidecar(out, ".cache.json"))

    glossary = dict(core.DEFAULT_GLOSSARY)
    if args.glossary:
        user_glossary = load_glossary(args.glossary)
        glossary.update(user_glossary)
        log.info("glossary: %d built-in term(s) plus %d from %s",
                 len(core.DEFAULT_GLOSSARY), len(user_glossary), args.glossary)
    else:
        log.info("glossary: %d built-in proper noun(s) locked to English "
                 "(use --glossary to override)", len(glossary))
    if args.no_glossary:
        glossary = {}
        log.info("glossary disabled by --no-glossary")

    prs = Presentation(str(src))
    insp = core.flatten(prs)
    parts = visible_parts(src)
    buckets: dict = {}
    part_source_skip = {u.source for u in insp.units}
    part_units = extract_part_units(src, buckets, part_source_skip) if parts else []
    if parts:
        log.info("chart/diagram parts with text: %d, extra strings: %d", len(buckets), len(part_units))

    all_units = insp.units + part_units
    if not args.notes:
        notes_only = [u for u in all_units if u.kind == "notes"]
        if notes_only:
            log.info("%d speaker-note paragraph(s) found; pass --notes to translate them", len(notes_only))

    todo_units = [u for u in all_units if u.kind != "notes" or args.notes]
    todo_units = [u for u in todo_units if not core.is_blank(u.text)]
    sources = [u.text for u in todo_units]

    log.info("slides: %d | translatable paragraphs: %d | unique strings: %d | characters: %d",
             insp.slides, len(todo_units), len(set(sources)), sum(len(s) for s in set(sources)))

    if args.dump_strings:
        Path(args.dump_strings).write_text(
            json.dumps(sorted(set(sources)), ensure_ascii=False, indent=1), "utf-8")
        log.info("source strings written to %s", args.dump_strings)

    report = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "input": str(src),
        "output": str(out),
        "slides": insp.slides,
        "paragraphs": len(todo_units),
        "unique_strings": len(set(sources)),
        "characters": sum(len(s) for s in set(sources)),
        "chart_parts": len(buckets),
        "pictures": len(insp.media),
        "provider": args.provider,
        # the real model, not the --model default: for local backends the model
        # name comes from --ollama-model / --local-model
        "model": (args.ollama_model if args.provider == "ollama"
                  else args.local_model or core.DEFAULT_LOCAL_MODEL if args.provider == "local"
                  else "mock" if args.provider == "mock"
                  else args.model or DEFAULT_MODEL),
    }

    if args.dry_run:
        preview = [
            {"slide": u.slide, "kind": u.kind, "where": u.where, "text": u.text}
            for u in todo_units[:40]
        ]
        report["mode"] = "dry-run"
        report["preview"] = preview
        _write_report(args.report or _sidecar(out, ".report.json"), report)
        _safe_print(json.dumps(report, ensure_ascii=False, indent=2)[:6000])
        return 0

    # Resolve the endpoint first: whether a key is required depends on it.
    if args.provider == "deepseek":
        base_url = args.base_url or DEFAULT_BASE_URL
        model = args.model or DEFAULT_MODEL
    elif args.provider == "openai":
        base_url = args.base_url or os.environ.get("OPENAI_BASE_URL", "https://api.openai.com/v1")
        model = args.model or os.environ.get("OPENAI_MODEL", "gpt-4o-mini")
    elif args.provider == "ollama":
        base_url = args.base_url or args.ollama_host
        model = args.model or args.ollama_model
    elif args.provider == "local":
        base_url = "local://in-process"
        model = args.local_model or core.DEFAULT_LOCAL_MODEL
    else:
        base_url = args.base_url or "mock://local"
        model = args.model or "mock"

    api_key, key_source = resolve_api_key(args.api_key)
    local_server = args.provider == "local" or is_local_endpoint(base_url)
    if args.provider != "mock" and not args.strings_json and not api_key and not local_server:
        log.error("no API key found: pass --api-key or set one of %s",
                  ", ".join(KEY_ENV_CANDIDATES))
        return 2
    if local_server:
        api_key = api_key or ""          # Ollama and friends need no credential
        log.info("local inference endpoint (%s): no API key required", base_url)
    elif key_source == "--api-key":
        log.warning("passing --api-key on the command line exposes it in shell history; prefer an env var")
    elif key_source:
        log.info("api key from %s", key_source)

    # A single local GPU/CPU serves one request at a time: fanning out only adds
    # queueing and memory pressure, and a small model is fastest without it.
    if args.provider in ("ollama", "local"):
        if args.concurrency > 1:
            log.info("--provider %s: forcing concurrency 1 (one local model, no fan-out)",
                     args.provider)
            args.concurrency = 1
        if args.batch_size > 1:
            args.batch_size = 1

    cfg = core.ProviderConfig(
        provider=args.provider,
        base_url=base_url,
        api_key=api_key or "",
        model=model,
        target_language=args.target_language,
        timeout=args.timeout,
        retries=args.retries,
        keep_terms=args.keep_terms,
    )
    # 'ollama' is an OpenAI-compatible local server, so it shares the HTTP provider.
    provider = core.build_provider(cfg)
    if args.provider == "local":
        ok, why = local.is_available()
        if not ok:
            log.error("--provider local is unavailable: %s", why)
            log.error("install it with:  pip install -e \".[local]\"")
            return 2
        provider = local.build_local_provider(
            cfg,
            model_id=args.local_model,
            device=args.local_device,
            prompt_style=args.local_prompt_style,
            cache_dir=args.local_cache_dir,
            max_new_tokens=args.local_max_new_tokens,
            local_files_only=args.local_files_only,
            target_language=_local_target(args.target_language),
        )
        log.info("local model: %s (device=%s, style=%s, endpoint=%s)",
                 provider.local.model_id, args.local_device, args.local_prompt_style,
                 local.hf_endpoint())

    if args.strings_json:
        mapping = json.loads(Path(args.strings_json).read_text("utf-8"))
        if not isinstance(mapping, dict):
            log.error("--strings-json must contain a JSON object of source -> translation")
            return 2
        resolved = {s: str(mapping.get(s, s)) for s in sources}
        translated = sum(1 for s in sources if resolved[s] != s)
        log.info("offline translation map: %d/%d string(s) supplied by %s",
                 translated, len(set(sources)), args.strings_json)
        report["provider"] = "strings-json"
        report["model"] = "n/a"
        report["translated_strings"] = translated
        report["untranslated_strings"] = len(set(sources)) - translated
    else:
        log.info("translating with %s (%s) at %s", args.provider, model, base_url)
        translator = core.Translator(
            provider,
            batch_size=args.batch_size,
            concurrency=args.concurrency,
            cache_path=cache_path,
            use_cache=not args.no_cache,
            glossary=glossary,
        )
        log.debug("cache: path=%s loaded=%d use_cache=%s digest=%s",
                  cache_path, len(translator.cache), not args.no_cache,
                  translator._glossary_digest)

        def progress(done: int, total: int) -> None:
            if total:
                log.info("progress: %d/%d strings (%.0f%%)", done, total, 100.0 * done / total)

        resolved = translator.translate_texts(sources, progress=progress)
        # Report only genuine failures. A string whose translation equals its
        # source is either a real failure or a deliberate keep-as-is (an acronym
        # the glossary pins), so it is never counted as missing on that basis.
        missing = sorted(translator.failed_sources)
        failures = sorted(set(getattr(provider, "failed_items", [])))
        if missing:
            log.warning("%d string(s) were not translated and stay in English", len(missing))
            log.warning("re-run the same command to retry only those (successes are cached)")
        report["translated_strings"] = len(set(sources)) - len(missing)
        report["untranslated_strings"] = len(missing)
        report["untranslated"] = missing[:200]
        report["api_failures"] = len(failures)

    # ---- write text frames back ------------------------------------------ #
    stats = {"frames": 0}
    apply_unitwise(prs, todo_units, resolved, args.notes, stats)
    log.info("updated %d text frame(s)", stats["frames"])

    # ---- chart / diagram parts ------------------------------------------- #
    part_updates: dict[str, bytes] = {}
    part_changed = 0
    for name, entry in buckets.items():
        changed = core.apply_part_text(buckets, name, resolved)
        if changed:
            part_changed += changed
            part_updates[name] = core.serialize_part(entry["root"])
    if part_updates:
        log.info("rewrote %d chart/diagram part(s), %d value(s)", len(part_updates), part_changed)
    report["chart_values_translated"] = part_changed

    # ---- optional auto-fit ------------------------------------------------ #
    if args.autofit != "off" or args.autofit_preview:
        try:
            include_titles = args.autofit == "all"
            fstats = fit.shrink_to_fit(prs, include_titles=include_titles,
                                       dry_run=args.autofit_preview)
            if args.autofit_preview:
                log.info("autofit preview: %d shape(s) would shrink (titles %s)",
                         len(fstats["would_shrink"]),
                         "included" if include_titles else "kept")
                report["autofit_preview"] = fstats["would_shrink"][:200]
            else:
                log.info("auto-fit: shrank %d fixed-size shape(s), %d table cell(s); "
                         "%d frame(s) left to PowerPoint autofit, %d within tolerance",
                         fstats["shapes"], fstats["cells"],
                         fstats.get("handled_by_powerpoint", 0), fstats["skip"])
                report["autofit"] = {k: v for k, v in fstats.items() if k != "would_shrink"}
        except Exception as exc:
            log.warning("auto-fit skipped: %s", exc)

    # ---- save, then splice the parts back in ----------------------------- #
    out.parent.mkdir(parents=True, exist_ok=True)
    if part_updates:
        with tempfile.TemporaryDirectory(prefix="pptx-translate-") as tmp:
            stage = Path(tmp) / "stage.pptx"
            prs.save(str(stage))
            core.rewrite_package(stage, out, part_updates)
    else:
        prs.save(str(out))

    # ---- verify ----------------------------------------------------------- #
    verification = verify_output(src, out)
    report["verification"] = verification
    report["images_with_text"] = _picture_note(insp)
    report["elapsed_seconds"] = round(time.time() - started, 1)

    if not args.quiet:
        _safe_print("")
        log.info("input : %s (%.1f MB)", src, src.stat().st_size / 1048576)
        log.info("output: %s (%.1f MB)", out, out.stat().st_size / 1048576)
        log.info("slides %s | pictures %s | media parts %s | animations %s",
                 verification["slides_ok"], verification["pictures"], verification["media"],
                 verification["timing_nodes"])

    if verification["slides"] != insp.slides:
        log.error("slide count changed: %d -> %d", insp.slides, verification["slides"])
        return 1

    _write_report(args.report or (out.with_suffix(out.suffix + ".report.json")), report)
    return 0


def _picture_note(insp: core.Inspection) -> list[dict]:
    """Pictures big enough to plausibly carry baked-in text."""
    out = []
    for item in insp.media:
        if item["width_in"] * item["height_in"] > 4.0:
            out.append(item)
    return out


def verify_output(src: Path, out: Path) -> dict:
    """Compare the package skeleton of input and output."""
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(out) as b:
        na, nb = a.namelist(), b.namelist()
        media_a = {n: len(a.read(n)) for n in na if n.startswith("ppt/media/")}
        media_b = {n: len(b.read(n)) for n in nb if n.startswith("ppt/media/")}
        anim = 0
        for n in nb:
            if n.endswith(".xml"):
                data = b.read(n)
                anim += data.count(b"<p:timing") + data.count(b"<p:anim")
    return {
        "slides": sum(1 for n in nb if n.startswith("ppt/slides/slide") and n.endswith(".xml")),
        "slides_ok": sum(1 for n in nb if n.startswith("ppt/slides/slide") and n.endswith(".xml")),
        "layouts": sum(1 for n in nb if n.startswith("ppt/slideLayouts/slideLayout")),
        "masters": sum(1 for n in nb if n.startswith("ppt/slideMasters/")),
        "pictures": sum(1 for n in nb if n.startswith("ppt/media/")),
        "media": len(media_b),
        "media_identical": media_a == media_b,
        "timing_nodes": anim,
        "parts_added": sorted(set(nb) - set(na))[:20],
        "parts_removed": sorted(set(na) - set(nb))[:20],
    }


def _write_report(path, report: dict) -> None:
    try:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(report, ensure_ascii=False, indent=2), "utf-8")
        log.info("report: %s", path)
    except Exception as exc:
        log.warning("could not write report: %s", exc)


if __name__ == "__main__":
    raise SystemExit(main())
