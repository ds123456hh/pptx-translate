# -*- coding: utf-8 -*-
"""Verify that a translated deck keeps every structural and style attribute.

Usage: python tools_verify.py ORIGINAL.pptx TRANSLATED.pptx
"""
from __future__ import annotations

import re
import sys
import zipfile
from pptx import Presentation


def para_texts(tf) -> list[str]:
    out = []
    for para in tf.paragraphs:
        chunks = []
        for child in para._p:
            tag = child.tag
            if tag.endswith("}r"):
                t = child.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t")
                if t is not None and t.text:
                    chunks.append(t.text)
            elif tag.endswith("}br"):
                chunks.append("\x0b")
        out.append("".join(chunks))
    return out


def shape_records(shapes, prefix=""):
    recs = {}
    for sh in shapes:
        key = f"{prefix}{sh.shape_id}:{sh.name}"
        try:
            rec = {
                "type": sh.shape_type,
                "left": int(sh.left), "top": int(sh.top),
                "width": int(sh.width), "height": int(sh.height),
                "paras": para_texts(sh.text_frame) if sh.has_text_frame else [],
                "runs": [],
            }
        except Exception as exc:
            recs[key] = {"error": str(exc)}
            continue
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    f = r.font
                    rec["runs"].append((
                        f.name, None if f.size is None else f.size.pt, f.bold, f.italic,
                        None if f.color is None or f.color.type is None else str(f.color.type),
                    ))
        if getattr(sh, "has_group", False) and sh.has_group:
            rec["children"] = shape_records(sh.shapes, prefix=key + "/")
        if getattr(sh, "has_table", False) and sh.has_table:
            cells = []
            for row in sh.table.rows:
                for cell in row.cells:
                    cells.append(para_texts(cell.text_frame))
            rec["table"] = cells
        recs[key] = rec
    return recs


def geometry(recs):
    out = {}
    for key, rec in recs.items():
        if "error" in rec:
            continue
        out[key] = (str(rec["type"]), rec["left"], rec["top"], rec["width"], rec["height"])
        if "children" in rec:
            out.update(geometry(rec["children"]))
    return out


def styles(recs):
    out = {}
    for key, rec in recs.items():
        if "error" in rec:
            continue
        out[key] = rec["runs"]
        if "children" in rec:
            out.update(styles(rec["children"]))
    return out


def main() -> int:
    src, dst = Presentation(sys.argv[1]), Presentation(sys.argv[2])
    a, b = list(src.slides), list(dst.slides)
    print(f"slides: {len(a)} -> {len(b)}", "OK" if len(a) == len(b) else "MISMATCH")

    geom_diff = 0
    style_diff = 0
    shape_diff = 0
    changed = same = 0
    empty_after = 0
    examples = []

    for idx, (sa, sb) in enumerate(zip(a, b), 1):
        ra, rb = shape_records(sa.shapes), shape_records(sb.shapes)
        ga, gb = geometry(ra), geometry(rb)
        if set(ga) != set(gb):
            shape_diff += 1
            if len(examples) < 6:
                examples.append(f"slide {idx}: shape set differs only-in-A={sorted(set(ga) - set(gb))[:3]} only-in-B={sorted(set(gb) - set(ga))[:3]}")
        for key in set(ga) & set(gb):
            if ga[key] != gb[key]:
                geom_diff += 1
                if len(examples) < 12:
                    examples.append(f"slide {idx} {key}: geometry {ga[key]} -> {gb[key]}")
        sa_, sb_ = styles(ra), styles(rb)
        for key in set(sa_) & set(sb_):
            if sa_[key] != sb_[key] and len(sa_[key]) == len(sb_[key]):
                style_diff += 1
                if len(examples) < 18:
                    examples.append(f"slide {idx} {key}: style {sa_[key]} -> {sb_[key]}")
        for key in set(ra) & set(rb):
            if "error" in ra[key] or "error" in rb[key]:
                continue
            for pa, pb in zip(ra[key]["paras"], rb[key]["paras"]):
                if not pa.strip():
                    continue
                if re.search(r"[A-Za-z]", pa):
                    if pa == pb:
                        same += 1
                        if len(examples) < 24:
                            examples.append(f"slide {idx} {key}: UNCHANGED {pa[:70]!r}")
                    else:
                        changed += 1
                        if not pb.strip():
                            empty_after += 1

    print(f"shape-set differences : {shape_diff}")
    print(f"geometry differences  : {geom_diff}")
    print(f"style differences     : {style_diff}")
    print(f"english paragraphs    : {changed} changed, {same} unchanged")
    print(f"emptied by translation: {empty_after}")
    for line in examples:
        print("   ", line)

    with zipfile.ZipFile(sys.argv[1]) as za, zipfile.ZipFile(sys.argv[2]) as zb:
        ma = {n: len(za.read(n)) for n in za.namelist() if n.startswith("ppt/media/")}
        mb = {n: len(zb.read(n)) for n in zb.namelist() if n.startswith("ppt/media/")}
        print("media identical:", ma == mb, f"({len(ma)} files)")
        ta = sum(za.read(n).count(b"<p:timing") for n in za.namelist() if n.endswith(".xml"))
        tb = sum(zb.read(n).count(b"<p:timing") for n in zb.namelist() if n.endswith(".xml"))
        print("p:timing nodes:", ta, "->", tb, "OK" if ta == tb else "MISMATCH")
        print("parts: %d -> %d" % (len(za.namelist()), len(zb.namelist())))
    ok = shape_diff == 0 and geom_diff == 0 and style_diff == 0 and same == 0 and len(a) == len(b)
    print("VERDICT:", "PASS" if ok else "REVIEW")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
